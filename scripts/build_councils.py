"""
Builds data/councils.js for councils.html: how does your council compare?

Victoria: the Local Government Performance Reporting Framework (Know Your
Council), 2020-21 to 2024-25, one spreadsheet published by the Victorian
Government. NSW: the Office of Local Government's "Your Council" time series,
one spreadsheet a year. Both are saved in raw/councils/ and refreshed by hand
once a year when the governments publish new figures.

Fairness rules built in:
  * councils are only compared with their official peer group, never ranked
    against the whole state, and never across states;
  * no overall score: each measure stands on its own;
  * figures are averaged over up to three years, to smooth one-off swings;
  * every measure says which way is better, or that there is no right answer;
  * a year whose typical value jumps implausibly against the other years is
    dropped (and reported), which catches mislabelled or rescaled columns.

    python scripts/build_councils.py
"""
from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "raw" / "councils"
OUT = ROOT / "data" / "councils.js"
MIN_PEERS = 5          # fewer councils with data than this: no comparison shown
JUMP = 1.8             # a year's median this many times off the others' is not trusted
STABLE = 0.5           # councils' ranks must hold up from year to year at least this well
SWING = 2.0            # a council's own range over the years, in multiples of the peer spread

# id, theme, everyday label, unit, better ("higher", "lower" or None), note
# units: pct (stored as a fraction), days, dollars, num, score (out of 100), ratio
VIC = [
    ("WC2", "Bins and recycling", "Bins missed, per 10,000 collections", "num", "lower", None),
    ("WC5", "Bins and recycling", "Kerbside waste kept out of landfill", "pct", "higher", None),
    ("WC3", "Bins and recycling", "Cost of collecting a garbage bin, a year", "dollars", "lower", None),
    ("SP1", "Planning permits", "Days to decide a planning application", "days", "lower", "Median days."),
    ("SP2", "Planning permits", "Planning decisions made on time", "pct", "higher", None),
    ("SP4", "Planning permits", "Decisions upheld when appealed at VCAT", "pct", "higher", None),
    ("SP3", "Planning permits", "Cost of the planning service, per application", "dollars", "lower", None),
    ("R2", "Local roads", "Sealed roads kept to the council's standard", "pct", "higher", None),
    ("R5", "Local roads", "Residents' satisfaction with local roads", "score", "higher", "Score out of 100, from the state's community survey."),
    ("R1", "Local roads", "Complaints about sealed roads, per 100 km", "num", "lower", None),
    ("LB8", "Libraries", "Library visits per resident", "num", "higher", None),
    ("LB6", "Libraries", "Library loans per resident", "num", "higher", None),
    ("LB5", "Libraries", "Library spending per resident", "dollars", None, "Lower is cheaper but can mean less service, so neither direction is better."),
    ("AF6", "Pools and leisure", "Pool visits per resident", "num", "higher", None),
    ("AF7", "Pools and leisure", "Net cost to the council per pool visit", "dollars", "lower", None),
    ("AM1", "Pets and animals", "Days to respond to an animal request", "days", "lower", None),
    ("AM5", "Pets and animals", "Collected animals rehomed", "pct", "higher", None),
    ("FS1", "Food safety", "Days to act on a food complaint", "days", "lower", None),
    ("FS4", "Food safety", "Serious food safety breaches followed up", "pct", "higher", None),
    ("MC4", "Maternal and child health", "Children taking part in maternal and child health checks", "pct", "higher", None),
    ("MC3", "Maternal and child health", "Cost of the service, per hour of nursing", "dollars", "lower", None),
    ("G5", "Decisions and trust", "Residents' satisfaction with council decisions", "score", "higher", "Score out of 100, from the state's community survey."),
    ("G2", "Decisions and trust", "Residents' satisfaction with community consultation", "score", "higher", "Score out of 100, from the state's community survey."),
    ("G1", "Decisions and trust", "Council decisions made behind closed doors", "pct", "lower", None),
    ("G3", "Decisions and trust", "Councillors' attendance at meetings", "pct", "higher", None),
    ("G4", "Decisions and trust", "Cost of elected councillors, per councillor", "dollars", "lower", None),
    ("OP1", "Money", "Underlying surplus, as a share of income", "pct", "higher", "Above zero means everyday income covers everyday costs."),
    ("L1", "Money", "Current assets for each dollar of bills due within a year", "ratio", "higher", None),
    ("O2", "Money", "Borrowings compared with rates", "pct", "lower", None),
    ("O5", "Money", "Spending on renewing assets compared with their wear", "pct", "higher", "100% means assets are renewed as fast as they wear out."),
    ("C7", "Money", "Staff turnover", "pct", "lower", None),
    ("E4", "Money", "Average rates per property", "dollars", None, "Depends heavily on property values and on what the council provides, so neither direction is better."),
    ("E2", "Money", "Spending per property", "dollars", None, "More spending can mean more services or less efficiency, so neither direction is better."),
]
VIC_CONTEXT = [("C6", "Socio-economic disadvantage decile", "num", "1 is the most disadvantaged tenth of Victoria, 10 the least."),
               ("C3", "Residents per km of local road", "num", None),
               ("C5", "Recurrent grants per resident", "dollars", None),
               ("C1", "Council spending per resident", "dollars", None)]

# NSW: column name prefix in the OLG spreadsheet
NSW = [
    ("da_days", "Planning permits", "Days to decide a development application", "days", "lower", "Mean gross days.", "Mean Gross Days for Development Applications"),
    ("recycling", "Bins and recycling", "Household waste recycled", "pct", "higher", None, "Recycling Rate"),
    ("waste_charge", "Bins and recycling", "Average domestic waste charge, a year", "dollars", None, "What the bin service costs households; it also reflects the services included, so neither direction is better.", "Average Domestic Waste Annual Charge"),
    ("asset_maint", "Roads and community assets", "Maintenance spent compared with what is needed", "pct", "higher", "100% means the council spends what its assets need.", "Asset Maintenance Ratio"),
    ("renewal", "Roads and community assets", "Spending on renewing assets compared with their wear", "pct", "higher", "100% means assets are renewed as fast as they wear out.", "Building & Infrastructure Renewal Ratio"),
    ("backlog", "Roads and community assets", "Backlog of assets in poor condition", "pct", "lower", "Share of asset value needing work to reach a satisfactory standard.", "Infrastructure Backlog Ratio"),
    ("library_loans", "Libraries", "Library loans per resident", "num", "higher", None, "Library Circulation Per Capita"),
    ("library_spend", "Libraries", "Library spending per resident", "dollars", None, "Lower is cheaper but can mean less service, so neither direction is better.", "Library Services Expenditure Per Capita"),
    ("pets", "Pets and animals", "Microchipped pets that are also registered", "pct", "higher", None, "Companion Animals Micro Chipped and Registered"),
    ("admin", "Running the council", "Governance and administration spending per resident", "dollars", "lower", None, "Governance & Administration Expenditure Per Capita"),
    ("staff", "Running the council", "Residents per council staff member", "num", None, "A leaner council or fewer services in-house, so neither direction is better.", "Population/Equivalent Full Time Staff"),
    ("op_perf", "Money", "Operating result, as a share of income", "pct", "higher", "Above zero means everyday income covers everyday costs.", "Operating Performance Ratio"),
    ("own_source", "Money", "Income raised by the council itself", "pct", "higher", "Rather than grants. The state benchmark is 60%.", "Own Source Revenue"),
    ("liquidity", "Money", "Spare cash for each dollar of bills due within a year", "ratio", "higher", "The state benchmark is 1.5.", "Unrestricted Current Ratio"),
    ("cash_cover", "Money", "Months of spending covered by cash", "num", "higher", "The state benchmark is 3 months.", "Cash Expense Cover Ratio"),
    ("outstanding", "Money", "Rates and charges unpaid at year end", "pct", "lower", None, "Rates and Annual Charges Outstanding"),
    ("res_rate", "Money", "Average residential rates", "dollars", None, "Depends heavily on land values and on what the council provides, so neither direction is better.", "Average Residential Rate"),
]
NSW_CONTEXT = [("Population 2", "Population", "num", None), ("Council Area", "Area, km²", "num", None),
               ("Population Density", "Residents per km²", "num", None),
               ("Population Change over 5 years", "Population growth over five years", "pct", None),
               ("Socio-Economic Index Rating", "Socio-economic index rating", "num", "A higher rating means less disadvantage.")]
NSW_FILES = [("2022-23", "nsw_2022_23.xlsx"), ("2023-24", "nsw_2023_24.xlsx"), ("2024-25", "nsw_2024_25.xlsx")]
VIC_GROUPS = {"metro": "Metropolitan", "interface": "Interface", "regional city": "Regional city",
              "large shire": "Large shire", "small shire": "Small shire"}
VIC_GROUP_TEXT = {"Metropolitan": "metropolitan Melbourne councils", "Interface": "interface councils on Melbourne's fringe",
                  "Regional city": "regional city councils", "Large shire": "large shire councils", "Small shire": "small shire councils"}


def num(x):
    if isinstance(x, (int, float)) and not isinstance(x, bool) and np.isfinite(x):
        return float(x)
    try:
        return float(str(x).strip())
    except (TypeError, ValueError):
        return None


def norm_cols(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


# measures that can be negative, where comparing years by ratio makes no sense
SIGNED = {"OP1", "op_perf"}
# plausible range for the typical (median) council, where only two years exist to compare
PLAUSIBLE = {"da_days": (20, 400)}


def trusted_years(values: dict[str, dict[str, float]], key: str, label: str, notes: list[str]) -> dict:
    """values: year -> {council: value}. Drops years that cannot be trusted: a median outside
    the plausible range, or (for measures that are never negative) a median far from the
    other years'. With only two years that disagree and no range to settle it, both go."""
    meds = {y: float(np.median(list(v.values()))) for y, v in values.items() if len(v) >= MIN_PEERS}
    if key in PLAUSIBLE:
        lo, hi = PLAUSIBLE[key]
        for y in [y for y, m in meds.items() if not lo <= m <= hi]:
            notes.append(f"{label}: {y} left out, its typical value ({meds[y]:.3g}) is outside the plausible range")
            del meds[y]
    if key not in SIGNED and len(meds) >= 2:
        far = lambda a, b: min(a, b) > 0 and max(a, b) / min(a, b) > JUMP  # noqa: E731
        if len(meds) == 2:
            a, b = meds.values()
            if far(a, b):
                notes.append(f"{label}: years disagree ({a:.3g} against {b:.3g}), measure left out")
                meds = {}
        else:
            for y in list(meds):
                ref = float(np.median([m for yy, m in meds.items() if yy != y]))
                if far(meds[y], ref):
                    notes.append(f"{label}: {y} left out, its typical value ({meds[y]:.3g}) is far from other years ({ref:.3g})")
                    del meds[y]
    return {y: values[y] for y in meds}


def stability(councils: list[dict], ind: str) -> float | None:
    """Average rank correlation between consecutive years across councils. A real
    measure of performance is fairly stable: councils that do well one year mostly
    do well the next. Low values mean the measure is mostly noise or timing."""
    years = sorted({y for c in councils for y in c["values"].get(ind, {}).get("years", {})})
    rs = []
    for a, b in zip(years, years[1:]):
        pairs = [(c["values"][ind]["years"][a], c["values"][ind]["years"][b]) for c in councils
                 if ind in c["values"] and a in c["values"][ind]["years"] and b in c["values"][ind]["years"]]
        if len(pairs) > 10:
            x, y = np.array(pairs).T
            rx, ry = np.argsort(np.argsort(x)), np.argsort(np.argsort(y))
            rs.append(float(np.corrcoef(rx, ry)[0, 1]))
    return float(np.mean(rs)) if rs else None


def screen(councils: list[dict], indicators: list[dict], state: str, notes: list[str]) -> tuple[list, list]:
    """Leaves out unstable measures, and flags councils whose own figures swing widely."""
    kept, left = [], []
    for ind in indicators:
        i = ind["id"]
        st = stability(councils, i)
        have = sum(i in c["values"] for c in councils)
        if have < MIN_PEERS * 2:
            left.append({**ind, "why": "Too few councils have reliable figures."})
        elif st is not None and st < STABLE:
            left.append({**ind, "why": f"Councils' positions change a lot from year to year (consistency {st:.2f} out of 1), "
                                       "so the figures say more about one-off events or accounting timing than performance."})
        else:
            kept.append({**ind, "stability": round(st, 2) if st is not None else None})
            continue
        notes.append(f"{state}, {ind['label']}: left out. {left[-1]['why']}")
        for c in councils:
            c["values"].pop(i, None)
    for ind in kept:
        i = ind["id"]
        for g in {c["group"] for c in councils}:
            cs = [c for c in councils if c["group"] == g and i in c["values"]]
            if len(cs) < MIN_PEERS:
                continue
            q1, q3 = np.percentile([c["values"][i]["v"] for c in cs], [25, 75])
            for c in cs:
                ys = list(c["values"][i]["years"].values())
                if len(ys) > 1 and q3 > q1 and (max(ys) - min(ys)) > SWING * (q3 - q1):
                    c["values"][i]["swings"] = True
    return kept, left


def compare(councils: list[dict], ind: str, better) -> None:
    """Adds each council's position within its peer group for one measure."""
    groups = {}
    for c in councils:
        if ind in c["values"]:
            groups.setdefault(c["group"], []).append(c)
    for g, cs in groups.items():
        if len(cs) < MIN_PEERS:
            for c in cs:
                c["values"][ind]["band"] = None
            continue
        vals = np.array([c["values"][ind]["v"] for c in cs])
        for c in cs:
            v = c["values"][ind]["v"]
            share = (np.sum(vals < v) + 0.5 * (np.sum(vals == v) - 1)) / (len(vals) - 1)   # 0 lowest, 1 highest
            if better == "lower":
                share = 1 - share
            c["values"][ind]["pos"] = round(float(share), 3)
            if better is None:
                band = "high" if share >= 0.75 else "low" if share <= 0.25 else "mid"
            else:
                band = "better" if share >= 0.75 else "worse" if share <= 0.25 else "typical"
            c["values"][ind]["band"] = band


# ---------------------------------------------------------------- Victoria
def build_vic(notes: list[str]) -> dict:
    wb = openpyxl.load_workbook(RAW / "vic_lgprf_2020_2025.xlsx", read_only=True, data_only=True)
    years = ["2022-23", "2023-24", "2024-25"]
    ids = {i[0] for i in VIC} | {c[0] for c in VIC_CONTEXT}
    raw, explain, group = {}, {}, {}
    for y in years:
        rows = list(wb[f"Indicators {y}"].iter_rows(values_only=True))
        why = [norm_cols(h) for h in rows[0]].index("Variation")    # the council's own explanation
        for r in rows[1:]:
            name, grp, iid = (r[0] or "").strip(), r[1], r[2]
            if not name or iid not in ids:
                continue
            if y == years[-1] and grp:
                group[name] = VIC_GROUPS[str(grp).strip().lower().replace("_", " ")]
            applicable = str(r[5] or "").lower().startswith("applicable")
            v = num(r[6])
            if applicable and v is not None:
                raw.setdefault(iid, {}).setdefault(y, {})[name] = v
                if y == years[-1]:
                    text = norm_cols(r[why])
                    # keep real explanations, not stock phrases like "Result is in the expected range"
                    if len(text) > 45 and not re.match(r"^(the )?results? (is|are) (with)?in (the )?expected range", text, re.I):
                        explain.setdefault(name, {})[iid] = text if len(text) <= 360 else text[:360].rsplit(" ", 1)[0] + "…"
    councils = [{"name": n, "group": g, "values": {}, "context": {}} for n, g in sorted(group.items())]
    by = {c["name"]: c for c in councils}
    for iid, theme, label, unit, better, note in VIC:
        yrs = trusted_years(raw.get(iid, {}), iid, f"Victoria, {label}", notes)
        for c in councils:
            ys = {y: yrs[y][c["name"]] for y in years if y in yrs and c["name"] in yrs[y]}
            if ys:
                c["values"][iid] = {"v": float(np.mean(list(ys.values()))), "years": {y: round(v, 4) for y, v in ys.items()}}
                if c["name"] in explain and iid in explain[c["name"]]:
                    c["values"][iid]["why"] = explain[c["name"]][iid]
        compare(councils, iid, better)
    for iid, label, unit, note in VIC_CONTEXT:
        latest = raw.get(iid, {}).get(years[-1], {})
        for n, v in latest.items():
            if n in by:
                by[n]["context"][iid] = v
    inds, left = screen(councils, [{"id": i, "theme": t, "label": l, "unit": u, "better": b, "note": n}
                                   for i, t, l, u, b, n in VIC], "Victoria", notes)
    return {
        "name": "Victoria", "left_out": left, "source": "Know Your Council, Local Government Performance Reporting Framework, Victorian Government",
        "years": "2022-23 to 2024-25", "peer_note": "Victoria's official council groups", "group_text": VIC_GROUP_TEXT,
        "indicators": inds,
        "context": [{"id": i, "label": l, "unit": u, "note": n} for i, l, u, n in VIC_CONTEXT],
        "councils": councils,
    }


# ---------------------------------------------------------------- NSW
def load_nsw(path: Path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows[:8]) if r[0] and str(r[0]).startswith("Council"))
    hdr = [norm_cols(c) for c in rows[h]]
    data = {str(r[0]).strip(): r for r in rows[h + 1:] if r[0] and num(r[1]) is not None}
    return hdr, data


def col(hdr, prefix):
    js = [j for j, c in enumerate(hdr) if c.lower().startswith(prefix.lower()) and "not collected" not in c.lower()]
    return js[0] if js else None


def build_nsw(notes: list[str]) -> dict:
    files = {y: load_nsw(RAW / f) for y, f in NSW_FILES}
    latest_y = NSW_FILES[-1][0]
    hdr, data = files[latest_y]
    gi, ci = 1, 2   # OLG group number, classification
    councils = [{"name": n, "group": str(r[ci]).strip(), "olg_group": int(num(r[gi])), "values": {}, "context": {}}
                for n, r in sorted(data.items())]
    by = {c["name"]: c for c in councils}
    for iid, theme, label, unit, better, note, prefix in NSW:
        per_year = {}
        for y, (h, d) in files.items():
            j = col(h, prefix)
            if j is None:
                continue
            vals = {n: num(r[j]) for n, r in d.items() if n in by and num(r[j]) is not None}
            if not vals:
                continue
            if unit == "pct" and np.median(list(vals.values())) > 1.5:     # this file stores 0-100
                vals = {n: v / 100 for n, v in vals.items()}
            per_year[y] = vals
        yrs = trusted_years(per_year, iid, f"NSW, {label}", notes)
        for c in councils:
            ys = {y: v[c["name"]] for y, v in yrs.items() if c["name"] in v}
            if ys:
                c["values"][iid] = {"v": float(np.mean(list(ys.values()))), "years": {y: round(v, 4) for y, v in ys.items()}}
        compare(councils, iid, better)
    for prefix, label, unit, note in NSW_CONTEXT:
        j = col(hdr, prefix)
        for n, r in data.items():
            v = num(r[j]) if j is not None else None
            if v is not None and n in by:
                if unit == "pct" and abs(v) > 1.5:
                    v /= 100
                by[n]["context"][prefix] = v
    groups = sorted({c["group"] for c in councils})
    inds, left = screen(councils, [{"id": i, "theme": t, "label": l, "unit": u, "better": b, "note": n}
                                   for i, t, l, u, b, n, _ in NSW], "NSW", notes)
    return {
        "name": "New South Wales", "left_out": left, "source": "Your Council time series data, Office of Local Government NSW",
        "years": "2022-23 to 2024-25", "peer_note": "the Office of Local Government's council classifications",
        "group_text": {g: g.lower() + " councils" for g in groups},
        "indicators": inds,
        "context": [{"id": p, "label": l, "unit": u, "note": n} for p, l, u, n in NSW_CONTEXT],
        "councils": councils,
    }


def main() -> None:
    notes: list[str] = []
    out = {"generated": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%MZ"),
           "states": {"vic": build_vic(notes), "nsw": build_nsw(notes)}, "notes": notes}
    for k, s in out["states"].items():
        cs = s["councils"]
        print(f"{s['name']}: {len(cs)} councils in {len({c['group'] for c in cs})} groups; "
              f"{sum(len(c['values']) for c in cs)} council measures")
    for n in notes:
        print("NOTE:", n)

    def tidy(o):
        if isinstance(o, float):
            return round(o, 4)
        if isinstance(o, dict):
            return {k: tidy(v) for k, v in o.items()}
        if isinstance(o, list):
            return [tidy(v) for v in o]
        return o
    OUT.write_text("window.COUNCILS = " + json.dumps(tidy(out), separators=(",", ":"), ensure_ascii=False) + ";\n",
                   encoding="utf-8")
    summary = {"councils": sum(len(s["councils"]) for s in out["states"].values()),
               "measures": sum(len(s["indicators"]) for s in out["states"].values())}
    (OUT.parent / "councils_summary.js").write_text("window.COUNCILS_SUMMARY = " + json.dumps(summary) + ";\n", encoding="utf-8")
    print(f"Wrote {OUT} ({OUT.stat().st_size / 1e3:.0f} kB) and the homepage summary")


if __name__ == "__main__":
    main()
