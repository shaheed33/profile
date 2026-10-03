"""
prepare_data.py
Turns Andrew Grogan's spreadsheet (sheet "20240930 NEM") into data/nem_data.js
for the website. Later, replace the reading part with NEMOSIS / Open Electricity
data -- as long as you write the same columns, the website doesn't change.

Usage:  python scripts/prepare_data.py "20240930 Open Electricity Simulation NL.xlsx"
Needs:  pip install openpyxl
"""
import sys, json, datetime as dt
import openpyxl

xlsx = sys.argv[1] if len(sys.argv) > 1 else "20240930 Open Electricity Simulation NL.xlsx"
wb = openpyxl.load_workbook(xlsx, data_only=True)
ws = wb["20240930 NEM"]

# Raw-sheet column letter -> name used by the website
COLS = {
    "D": "battery_charging", "E": "phes_pumps",
    "I": "coal_brown", "J": "coal_black",
    "K": "bio_biogas", "L": "bio_biomass",
    "M": "distillate", "N": "gas_steam", "O": "gas_ccgt", "P": "gas_ocgt",
    "Q": "gas_recip", "R": "gas_wcmg",
    "S": "battery_discharging", "U": "hydro", "V": "wind",
    "W": "solar_utility", "X": "solar_rooftop", "Z": "price",
}
SUPPLY = [c for c in COLS if c not in ("D", "E", "Z")]

data = {name: [] for name in COLS.values()}
data["demand"] = []
first_time = None
row = 9
while ws[f"B{row}"].value is not None and ws[f"I{row}"].value is not None:
    if first_time is None:
        # The raw sheet's clock is 10 hours ahead; the scaled sheet shows the
        # correct local time (AEST), so shift back 10 hours.
        first_time = ws[f"B{row}"].value - dt.timedelta(hours=10)
    vals = {c: float(ws[f"{c}{row}"].value or 0) for c in COLS}
    for c, name in COLS.items():
        data[name].append(round(vals[c], 2))
    # Native demand = generation minus storage charging (same as Andrew's column H)
    data["demand"].append(round(sum(vals[c] for c in SUPPLY) + vals["D"] + vals["E"], 2))
    row += 1

out = {
    "label": "29 Sep – 29 Oct 2024 (NEM, half-hourly)",
    "source": "Open Electricity, via Andrew Grogan's simulation workbook",
    "start": first_time.strftime("%Y-%m-%dT%H:%M"),
    "step_minutes": 30,
    "existing_storage_mwh": 10635,
    "series": data,
}
with open("data/nem_data.js", "w") as f:
    f.write("window.NEM_DATA = " + json.dumps(out, separators=(",", ":")) + ";\n")
print(f"Wrote {len(data['demand'])} half-hours starting {out['start']}")
